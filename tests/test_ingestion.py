import docx
import pytest
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from groundon.ingestion.chunker import chunk_text
from groundon.ingestion.parsers import UnsupportedFileType, parse_file


def test_chunks_respect_size_and_overlap():
    text = " ".join(f"Sentence number {i} is here." for i in range(60))
    chunks = chunk_text(text, chunk_size=120, chunk_overlap=40)
    assert len(chunks) > 5
    assert all(len(c) <= 120 for c in chunks)
    # Overlap: each chunk starts with the last sentence of the previous one.
    for prev, nxt in zip(chunks, chunks[1:]):
        last_sentence = prev.split(". ")[-1]
        assert nxt.startswith(last_sentence.rstrip(".")), (prev, nxt)


def test_overlong_sentence_is_hard_split():
    chunks = chunk_text("x" * 1000, chunk_size=300, chunk_overlap=50)
    assert all(len(c) <= 300 for c in chunks)
    assert "".join(chunks).count("x") >= 1000


def test_invalid_overlap_rejected():
    with pytest.raises(ValueError):
        chunk_text("hello", chunk_size=100, chunk_overlap=100)


def test_parse_pdf_keeps_page_numbers(tmp_path):
    path = tmp_path / "a.pdf"
    c = canvas.Canvas(str(path), pagesize=A4)
    c.drawString(72, 750, "First page text")
    c.showPage()
    c.drawString(72, 750, "Second page text")
    c.save()
    pages = parse_file(path)
    assert [p.page for p in pages] == [1, 2]
    assert "Second page" in pages[1].text


def test_parse_docx_includes_tables(tmp_path):
    path = tmp_path / "a.docx"
    d = docx.Document()
    d.add_paragraph("Intro paragraph")
    table = d.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Germany", "EUR 75"
    d.save(str(path))
    text = parse_file(path)[0].text
    assert "Intro paragraph" in text and "Germany | EUR 75" in text


def test_unsupported_type(tmp_path):
    path = tmp_path / "a.xlsx"
    path.write_bytes(b"x")
    with pytest.raises(UnsupportedFileType):
        parse_file(path)


def test_ingest_is_idempotent_and_deletable(make_services, docs_dir):
    s = make_services()
    assert {d["source"] for d in s.store.list_documents()} == {"leave.md", "travel.md", "security.txt"}
    size = s.store.size
    again = s.ingestor.ingest_file(docs_dir / "leave.md")
    assert again.status == "duplicate" and s.store.size == size

    doc_id = next(d["doc_id"] for d in s.store.list_documents() if d["source"] == "leave.md")
    assert s.store.delete_document(doc_id)
    assert s.store.size < size
    assert all(c.source != "leave.md" for c in s.store.all_chunks())


def test_index_persists_across_restarts(make_services):
    first = make_services()
    size = first.store.size
    second = make_services()  # reloads FAISS + SQLite from disk; auto-ingest skipped
    assert second.store.size == size
    assert second.retriever.retrieve("per diem Germany", top_k=1).chunks[0].chunk.source == "travel.md"


def test_overlap_never_pushes_chunk_over_size():
    # A short sentence carried as overlap plus a near-limit sentence used to exceed chunk_size.
    text = "A" * 90 + ". " + "B" * 480 + ". " + " ".join(["C" * 30 + "."] * 3)
    chunks = chunk_text(text, chunk_size=500, chunk_overlap=100)
    assert all(len(c) <= 500 for c in chunks), [len(c) for c in chunks]
    assert any("B" * 480 in c for c in chunks)
