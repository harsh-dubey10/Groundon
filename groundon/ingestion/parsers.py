"""Turn uploaded files into plain text, keeping page numbers where the format has them."""

from dataclasses import dataclass
from pathlib import Path

import docx
from pypdf import PdfReader

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}


@dataclass
class Page:
    text: str
    page: int | None  # 1-based for PDFs; None for formats without pages


class UnsupportedFileType(ValueError):
    pass


def parse_file(path: Path) -> list[Page]:
    ext = path.suffix.lower()
    if ext == ".pdf":
        return _parse_pdf(path)
    if ext == ".docx":
        return _parse_docx(path)
    if ext in {".txt", ".md"}:
        return [Page(text=path.read_text(encoding="utf-8", errors="replace"), page=None)]
    raise UnsupportedFileType(f"Unsupported file type '{ext}'. Supported: {sorted(SUPPORTED_EXTENSIONS)}")


def _parse_pdf(path: Path) -> list[Page]:
    reader = PdfReader(path)
    pages = []
    for number, pdf_page in enumerate(reader.pages, start=1):
        text = pdf_page.extract_text() or ""
        if text.strip():
            pages.append(Page(text=text, page=number))
    return pages


def _parse_docx(path: Path) -> list[Page]:
    document = docx.Document(str(path))
    # Blank line between paragraphs so the chunker can split on paragraph boundaries.
    blocks = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            blocks.append(" | ".join(cell.text.strip() for cell in row.cells))
    return [Page(text="\n\n".join(blocks), page=None)]
